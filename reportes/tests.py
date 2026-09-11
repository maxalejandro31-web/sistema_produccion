"""Tests de reportes: que "Cobros por Estancia" no incluya material propio
de la maquila, y que un valor numérico en 0 (rendimiento, scrap, merma,
peso restante) se muestre como "0" en el Excel en vez de una celda vacía
indistinguible de "sin dato".
"""
import datetime
import io

import openpyxl
from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from core.factories import crear_cliente, crear_mp, crear_linea

from produccion.models import OrdenProduccion


def _leer_excel(response):
    return openpyxl.load_workbook(io.BytesIO(response.content))


class CobrosEstanciaExcluyeMaterialPropioTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('reportes_user', password='pass12345')
        self.client.force_login(self.user)

    def test_material_propio_no_aparece_en_cobros_estancia(self):
        propia = crear_cliente(nombre='MAQUILAS Y SERVICIOS JC')
        cliente_externo = crear_cliente(nombre='Cliente Externo SA')
        hace_40_dias = timezone.localdate() - datetime.timedelta(days=40)

        mp_propia = crear_mp(numero_mp='MP-PROPIA-1', cliente=propia, fecha_entrada=hace_40_dias)
        mp_cliente = crear_mp(numero_mp='MP-CLIENTE-1', cliente=cliente_externo, fecha_entrada=hace_40_dias)

        resp = self.client.get(reverse('reporte_cobros_estancia'))
        wb = _leer_excel(resp)
        ws = wb.active
        numeros_mp = [row[0].value for row in ws.iter_rows(min_row=2)]

        self.assertNotIn('MP-PROPIA-1', numeros_mp)
        self.assertIn('MP-CLIENTE-1', numeros_mp)


class CerosRealesNoSalenEnBlancoTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('reportes_user2', password='pass12345')
        self.client.force_login(self.user)

    def test_rendimiento_cero_se_muestra_como_cero(self):
        cliente = crear_cliente()
        mp = crear_mp(cliente=cliente, peso=1000)
        linea = crear_linea()
        orden = OrdenProduccion.objects.create(
            tipo_proceso='corte_liso', mp=mp, cliente=cliente, linea=linea,
            peso_usado=100, peso_producido=0, estado='terminado',
        )

        resp = self.client.get(reverse('reporte_ordenes_produccion'))
        wb = _leer_excel(resp)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        col_rendimiento = headers.index('Rendimiento (%)')
        col_folio = headers.index('Folio')

        fila = next(row for row in ws.iter_rows(min_row=2) if row[col_folio].value == orden.folio_orden)
        self.assertEqual(fila[col_rendimiento].value, 0.0)

    def test_peso_restante_cero_se_muestra_como_cero_en_inventario_mp(self):
        mp = crear_mp(peso=100, peso_restante=0, estado='Terminado')

        resp = self.client.get(reverse('reporte_inventario_mp'))
        wb = _leer_excel(resp)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        col_peso_restante = headers.index('Peso Restante (kg)')
        col_numero = headers.index('N° MP')

        fila = next(row for row in ws.iter_rows(min_row=2) if row[col_numero].value == mp.numero_mp)
        self.assertEqual(fila[col_peso_restante].value, 0.0)
